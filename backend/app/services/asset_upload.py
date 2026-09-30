"""자산 엑셀 일괄 업로드 — specs/ASSET-1 D50~D52, P10.

기존에 엑셀로 관리하던 자산을 이관한다. 흐름: 템플릿 내려받기 → 채워서 [미리보기](검증만) → [등록].
서버는 상태를 들고 있지 않는다(Vercel 서버리스) — [등록] 때 같은 파일을 다시 받아 **처음부터 다시 검증**한다.

이 모듈은 DB 에 의존하지 않는다(파싱·검증·템플릿). 팀원 명단·기존 자산은 호출자가 넘겨 준다.
저장(채번 + 자산 + 배정 이력, 한 트랜잭션)은 `asset_registration.upload_assets`.

- 시트 = 유형(SW / HW / 교육), 1행 = 머리글, 2행부터 한 행 = 자산 1건(수량 열 없음 — D51).
- 열은 **머리글 이름**으로 찾는다(순서 무관, `*`·공백 무시). 그룹·소분류는 한글 이름(코드도 허용).
- 오류가 한 건이라도 있으면 전부 거부(D50). 중복 의심·자동 채움은 경고만(D52).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from ..config import (
    ASSET_HIDDEN_STATUSES,
    ASSET_MENU_CATEGORY_LABELS,
    ASSET_SUBCATEGORIES,
    ASSET_UPLOAD_MAX_ROWS,
    GROUPS,
)
from ..errors import AppError, validation
from ..models import Member
from ..schemas import AssetFieldsIn, AssetUploadIssue, AssetUploadPreview, AssetUploadRow
from .asset_bulk import check_valid_range, clean_fields


# --------------------------------------------------------------------------- 열 정의(템플릿 · 파싱 공용)
@dataclass(frozen=True)
class Col:
    header: str
    key: str
    kind: str = "text"  # text | date | int | group | sub
    required: bool = False
    width: int = 14


_HEAD = [
    Col("품명", "name", required=True, width=28),
    Col("소분류", "subcategory", "sub", width=12),
    Col("그룹", "group_code", "group", required=True, width=18),
    Col("구매일", "purchase_date", "date", width=12),
    Col("금액", "price", "int", width=12),
    Col("구매처", "purchased_from", width=16),
    Col("유효기간 시작", "valid_from", "date", width=13),
    Col("유효기간 종료", "valid_to", "date", width=13),
]
_TYPED = {
    "SW": [Col("버전", "version", width=10), Col("라이선스키", "license_key", width=26),
           Col("계정 ID", "account_id", width=20), Col("비밀번호", "password", width=14)],
    "HW": [Col("제조사", "manufacturer", width=12), Col("모델명", "model", width=16),
           Col("시리얼", "serial_no", width=18), Col("MAC", "mac_address", width=18)],
    "EDU": [Col("강의명", "course_title", width=28), Col("강의 URL", "course_url", width=30),
            Col("계정 ID", "account_id", width=20), Col("비밀번호", "password", width=14)],
}
_USE = [
    Col("사용자 사번", "member_id", width=12),
    Col("사용자 이름", "member_name", width=12),
    Col("공용 장소/용도", "shared_label", width=16),
    Col("사용 시작일", "start_date", "date", width=12),
]
_TAIL = [
    Col("견적번호", "quote_no", width=12),
    Col("계약번호", "contract_no", width=14),
    Col("기존 관리번호", "legacy_no", width=14),
    Col("비고", "note", width=30),
]
# 자산 문서 필드가 아닌 입력(배정·기존 번호) — 저장 전에 떼어 낸다
_NON_ASSET_KEYS = {"member_id", "member_name", "shared_label", "start_date", "legacy_no"}
# 중복 의심 경고(D52) 대상
_DUP_KEYS = [("serial_no", "시리얼"), ("license_key", "라이선스키"), ("account_id", "계정 ID")]
_TEMPLATE_ROWS = ASSET_UPLOAD_MAX_ROWS  # 드롭다운·서식을 미리 걸어 둘 행 수


def columns(category: str) -> list[Col]:
    cols = _HEAD + _TYPED[category] + _USE + _TAIL
    if not ASSET_SUBCATEGORIES[category]:
        cols = [c for c in cols if c.kind != "sub"]
    if category == "EDU":  # 교육 자산은 공용 배정 불가(D33)
        cols = [c for c in cols if c.key != "shared_label"]
    return cols


def sheet_name(category: str) -> str:
    return ASSET_MENU_CATEGORY_LABELS[category]


def _norm(s) -> str:
    return re.sub(r"[\s*]", "", str(s or ""))


def _header_text(c: Col) -> str:
    return f"{c.header} *" if c.required else c.header


# --------------------------------------------------------------------------- 템플릿
_GUIDE = [
    "자산 엑셀 일괄 업로드 양식",
    "",
    "1. SW / HW / 교육 시트에 한 행에 자산 1건씩 입력합니다. 쓰지 않는 시트는 비워 두세요.",
    "2. * 표시 열(품명, 그룹)은 필수입니다. 머리글(1행)은 지우거나 바꾸지 마세요.",
    "3. 그룹·소분류는 셀의 드롭다운에서 고릅니다.",
    "4. 날짜는 2026-01-31 형식(엑셀 날짜 셀도 가능), 금액은 원 단위 숫자입니다.",
    "5. 자산 번호는 시스템이 새로 붙입니다(구매일 연도 기준). 예전 번호는 '기존 관리번호'에 적으면 비고에 남습니다.",
    "6. 사용 중인 자산은 '사용자 사번' 또는 '사용자 이름'(또는 '공용 장소/용도')을 적습니다.",
    "   - 팀원 명단에 먼저 등록된 사람만 가능합니다. 이름만 적으면 재직 중인 팀원 중 한 명만 일치해야 합니다.",
    "   - '사용 시작일'을 비우면 구매일, 구매일도 없으면 업로드한 날로 채웁니다.",
    f"7. 파일당 최대 {ASSET_UPLOAD_MAX_ROWS}행. 오류가 한 건이라도 있으면 아무것도 등록되지 않습니다.",
    "8. 폐기된 자산은 올리지 마세요(살아 있는 자산만 이관).",
    "",
    "※ 비밀번호 열에 값을 넣었다면, 업로드가 끝난 뒤 이 파일을 반드시 삭제하세요.",
]
_REQ_FILL = PatternFill("solid", fgColor="FFF2CC")
_HEAD_FILL = PatternFill("solid", fgColor="E7E6E6")


def _list_validation(values: list[str]) -> DataValidation:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error = "목록에서 고르세요."
    dv.showErrorMessage = True
    return dv


def build_template_xlsx() -> bytes:
    wb = Workbook()
    guide = wb.active
    guide.title = "안내"
    guide.column_dimensions["A"].width = 110
    for r, line in enumerate(_GUIDE, start=1):
        cell = guide.cell(row=r, column=1, value=line)
        if r == 1:
            cell.font = Font(bold=True, size=14)
        elif line.startswith("※"):
            cell.font = Font(bold=True, color="C00000")

    last = _TEMPLATE_ROWS + 1
    for category in ASSET_MENU_CATEGORY_LABELS:
        ws = wb.create_sheet(sheet_name(category))
        ws.freeze_panes = "A2"
        for i, c in enumerate(columns(category), start=1):
            letter = get_column_letter(i)
            h = ws.cell(row=1, column=i, value=_header_text(c))
            h.font = Font(bold=True)
            h.fill = _REQ_FILL if c.required else _HEAD_FILL
            h.alignment = Alignment(horizontal="center")
            ws.column_dimensions[letter].width = c.width
            rng = f"{letter}2:{letter}{last}"
            if c.kind == "group":
                dv = _list_validation(list(GROUPS.values()))
            elif c.kind == "sub":
                dv = _list_validation(list(ASSET_SUBCATEGORIES[category].values()))
            else:
                dv = None
                # 사번·시리얼·버전(1.10) 등이 숫자로 바뀌어 앞자리 0·끝자리 0 이 사라지지 않게 텍스트 서식
                fmt = {"date": "yyyy-mm-dd", "int": "#,##0"}.get(c.kind, "@")
                for row in ws.iter_rows(min_row=2, max_row=last, min_col=i, max_col=i):
                    row[0].number_format = fmt
            if dv is not None:
                ws.add_data_validation(dv)
                dv.add(rng)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_filename() -> str:
    return "asset_upload_template.xlsx"


# --------------------------------------------------------------------------- 파싱
@dataclass
class RawRow:
    sheet: str
    row: int  # 엑셀 행 번호(머리글 = 1)
    category: str
    values: dict  # Col.key → 셀 원본 값


@dataclass
class Parsed:
    rows: list[RawRow]
    warnings: list[AssetUploadIssue] = field(default_factory=list)


def _sheet_category(title: str) -> str | None:
    t = _norm(title).upper()
    for cat, label in ASSET_MENU_CATEGORY_LABELS.items():
        if t in (cat, _norm(label).upper()):
            return cat
    return None


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def parse_workbook(content: bytes) -> Parsed:
    """파일 단위 문제(엑셀 아님, 시트 없음, 필수 머리글 없음, 행 없음, 상한 초과)는 400 으로 바로 거부."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise validation("엑셀(.xlsx) 파일을 읽을 수 없습니다. 내려받은 템플릿에 입력해 올려 주세요.")

    parsed = Parsed(rows=[])
    found = False
    try:
        for ws in wb.worksheets:
            category = _sheet_category(ws.title)
            if category is None:
                continue
            found = True
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if header is None:
                continue
            by_norm = {_norm(c.header): c for c in columns(category)}
            index: dict[int, Col] = {}
            for i, h in enumerate(header):
                if _blank(h):
                    continue
                col = by_norm.get(_norm(h))
                if col is None:
                    parsed.warnings.append(AssetUploadIssue(
                        sheet=ws.title, column=str(h), message="알 수 없는 열이라 무시합니다."))
                else:
                    index[i] = col
            rows = []
            for r, values in enumerate(it, start=2):
                vals = {col.key: values[i] for i, col in index.items() if i < len(values)}
                if all(_blank(v) for v in vals.values()):
                    continue
                rows.append(RawRow(ws.title, r, category, vals))
            if rows:
                missing = [c.header for c in columns(category) if c.required and c not in index.values()]
                if missing:
                    raise validation(f"'{ws.title}' 시트에 필수 열이 없습니다: {', '.join(missing)}. 템플릿 머리글을 그대로 쓰세요.")
            parsed.rows.extend(rows)
            if len(parsed.rows) > ASSET_UPLOAD_MAX_ROWS:
                raise validation(f"한 번에 {ASSET_UPLOAD_MAX_ROWS}행까지 올릴 수 있습니다. 파일을 나눠서 올려 주세요.")
    finally:
        wb.close()

    if not found:
        names = "/".join(ASSET_MENU_CATEGORY_LABELS.values())
        raise validation(f"{names} 시트를 찾을 수 없습니다. 내려받은 템플릿을 사용하세요.")
    if not parsed.rows:
        raise validation("등록할 행이 없습니다.")
    return parsed


# --------------------------------------------------------------------------- 셀 값 변환
_DATE_RE = re.compile(r"^(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})\.?$")


def _to_text(v) -> str | None:
    if _blank(v):
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _to_date(v) -> str | None:
    if _blank(v):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 1 <= v < 2958466:  # 엑셀 날짜 일련번호
        return from_excel(v).date().isoformat()
    m = _DATE_RE.match(str(v).strip())
    if m:
        try:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        except ValueError:
            pass
    raise ValueError("날짜 형식이 올바르지 않습니다(예: 2026-01-31).")


def _to_int(v) -> int | None:
    if _blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError("금액은 숫자로 입력하세요.")
    if isinstance(v, str):
        s = re.sub(r"[,\s원₩]", "", v)
        if not re.fullmatch(r"-?\d+(\.0+)?", s):
            raise ValueError("금액은 숫자로 입력하세요.")
        v = float(s)
    if isinstance(v, float):
        if not v.is_integer():
            raise ValueError("금액은 원 단위 정수로 입력하세요.")
        v = int(v)
    if v < 0:
        raise ValueError("금액은 0 이상이어야 합니다.")
    return v


def _match_label(v: str, options: dict[str, str], what: str) -> str:
    """코드 또는 한글 이름 → 코드."""
    n = _norm(v).upper()
    for code, label in options.items():
        if n in (code.upper(), _norm(label).upper()):
            return code
    raise ValueError(f"{what}은(는) {', '.join(options.values())} 중 하나여야 합니다.")


def _coerce(c: Col, v, category: str):
    if c.kind == "date":
        return _to_date(v)
    if c.kind == "int":
        return _to_int(v)
    s = _to_text(v)
    if s is None:
        return None
    if c.kind == "group":
        return _match_label(s, GROUPS, "그룹")
    if c.kind == "sub":
        return _match_label(s, ASSET_SUBCATEGORIES[category], "소분류")
    return s


# --------------------------------------------------------------------------- 검증
@dataclass(frozen=True)
class Target:
    """배정 대상 — 팀원(member_id) 또는 공용(shared_label) 중 하나."""

    member_id: str | None
    member_name: str | None
    shared_label: str | None
    scope_group_code: str


@dataclass
class PlannedAsset:
    category: str
    fields: dict  # clean_fields 결과(비밀번호는 아직 평문 — 저장 직전에 암호화)
    target: Target | None
    start_date: str | None


@dataclass
class Plan:
    items: list[PlannedAsset]
    preview: AssetUploadPreview

    @property
    def ok(self) -> bool:
        return not self.preview.errors


def _headers(category: str) -> dict[str, str]:
    return {c.key: c.header for c in columns(category)}


def build_plan(
    parsed: Parsed,
    members: list[Member],
    existing_assets: list[dict],
    *,
    today: date,
) -> Plan:
    errors: list[AssetUploadIssue] = []
    warnings: list[AssetUploadIssue] = list(parsed.warnings)
    items: list[PlannedAsset] = []
    rows: list[AssetUploadRow] = []
    counts = {c: 0 for c in ASSET_MENU_CATEGORY_LABELS}

    by_id = {m.employee_no: m for m in members}
    active_by_name: dict[str, list[Member]] = {}
    for m in members:
        if m.active:
            active_by_name.setdefault(m.name, []).append(m)
    unknown_people: list[str] = []

    live = [a for a in existing_assets if a.get("status") not in ASSET_HIDDEN_STATUSES]
    existing_dup = {k: {} for k, _ in _DUP_KEYS}
    for a in live:
        for k, _ in _DUP_KEYS:
            if a.get(k):
                existing_dup[k].setdefault(a[k], a.get("asset_no"))
    seen_dup = {k: {} for k, _ in _DUP_KEYS}
    today_s = today.isoformat()

    for raw in parsed.rows:
        cat = raw.category
        heads = _headers(cat)
        row_errors: list[AssetUploadIssue] = []

        def err(key: str | None, msg: str) -> None:
            row_errors.append(AssetUploadIssue(sheet=raw.sheet, row=raw.row, column=heads.get(key) if key else None, message=msg))

        def warn(key: str | None, msg: str) -> None:
            warnings.append(AssetUploadIssue(sheet=raw.sheet, row=raw.row, column=heads.get(key) if key else None, message=msg))

        vals: dict = {}
        for c in columns(cat):
            try:
                vals[c.key] = _coerce(c, raw.values.get(c.key), cat)
            except ValueError as e:
                vals[c.key] = None
                err(c.key, str(e))
        for c in columns(cat):
            if c.required and vals.get(c.key) is None and not any(e.column == c.header for e in row_errors):
                err(c.key, f"{c.header}은(는) 필수입니다.")

        # 자산 필드 정리 — 기존 관리번호는 비고 맨 앞에 남긴다
        data = {k: v for k, v in vals.items() if k not in _NON_ASSET_KEYS}
        if vals.get("legacy_no"):
            data["note"] = f"기존 관리번호: {vals['legacy_no']}" + (f"\n{data['note']}" if data.get("note") else "")
        fields: dict = {}
        try:
            AssetFieldsIn.model_validate(data)  # 길이 제한 등(등록 화면과 같은 규칙)
        except ValidationError as e:
            for x in e.errors():
                key = str(x["loc"][0]) if x["loc"] else None
                err(key, "입력값이 너무 깁니다." if x["type"] == "string_too_long" else x["msg"])
        if not row_errors:  # 필수값 누락 등을 같은 행에서 두 번 알리지 않게
            try:
                fields = clean_fields(cat, data)
                check_valid_range(fields.get("valid_from"), fields.get("valid_to"))
            except AppError as e:
                err("valid_to" if "유효기간" in e.message else None, e.message)

        # 배정 대상(D50 — 사번 우선, 없으면 이름)
        target: Target | None = None
        mid, mname, shared = vals.get("member_id"), vals.get("member_name"), vals.get("shared_label")
        if shared and (mid or mname):
            err("shared_label", "팀원과 공용(장소/용도) 중 하나만 입력하세요.")
        elif mid:
            m = by_id.get(mid)
            if m is None:
                err("member_id", f"팀원 명단에 없는 사번입니다: {mid}")
                unknown_people.append(f"{mname or ''}({mid})" if mname else mid)
            elif not m.active:
                err("member_id", f"퇴사 처리된 팀원({m.name})에게는 배정할 수 없습니다.")
            else:
                if mname and mname != m.name:
                    warn("member_name", f"사번 {mid}의 명단 이름은 '{m.name}'입니다. 사번 기준으로 배정합니다.")
                target = Target(m.employee_no, m.name, None, m.group_code)
        elif mname:
            cands = active_by_name.get(mname, [])
            if not cands:
                err("member_name", f"재직 중인 팀원 명단에 없는 이름입니다: {mname}")
                unknown_people.append(mname)
            elif len(cands) > 1:
                nos = ", ".join(m.employee_no for m in cands)
                err("member_name", f"같은 이름의 팀원이 여러 명입니다({nos}). 사용자 사번을 입력하세요.")
            else:
                m = cands[0]
                target = Target(m.employee_no, m.name, None, m.group_code)
        elif shared and vals.get("group_code"):
            target = Target(None, None, shared, vals["group_code"])  # 공용은 등록 그룹 스코프(D33)

        start = vals.get("start_date")
        start_auto = False
        if target is not None:
            if not start:
                start = vals.get("purchase_date") or today_s
                start_auto = True
                warn("start_date", f"사용 시작일이 비어 있어 {start}({'구매일' if vals.get('purchase_date') else '오늘'})로 채웁니다.")
            if start > today_s:
                err("start_date", "사용 시작일은 미래 날짜일 수 없습니다(오늘까지만 가능).")
            elif vals.get("purchase_date") and start < vals["purchase_date"]:
                err("start_date", f"사용 시작일은 구매일({vals['purchase_date']})보다 앞설 수 없습니다.")
        elif start and not row_errors:
            warn("start_date", "사용자가 없어 사용 시작일은 무시합니다(미사용으로 등록).")
            start = None

        # 중복 의심(D52) — 경고만
        for k, label in _DUP_KEYS:
            v = fields.get(k)
            if not v:
                continue
            if v in existing_dup[k]:
                warn(k, f"{label}이(가) 이미 등록된 자산 {existing_dup[k][v]}과 같습니다.")
            if v in seen_dup[k]:
                s, r = seen_dup[k][v]
                warn(k, f"{label}이(가) {s} 시트 {r}행과 같습니다.")
            else:
                seen_dup[k][v] = (raw.sheet, raw.row)

        errors.extend(row_errors)
        counts[cat] += 1
        sub = vals.get("subcategory")
        rows.append(AssetUploadRow(
            sheet=raw.sheet, row=raw.row, category=cat,
            name=vals.get("name"),
            group_name=GROUPS.get(vals.get("group_code") or ""),
            subcategory_label=ASSET_SUBCATEGORIES[cat].get(sub) if sub else None,
            user_label=(f"{target.member_name}({target.member_id})" if target and target.member_id
                        else f"공용 · {target.shared_label}" if target else None),
            start_date=start if target else None,
            start_auto=start_auto,
            has_password=bool(fields.get("password")),
        ))
        if not row_errors:
            items.append(PlannedAsset(cat, fields, target, start if target else None))

    if unknown_people:
        uniq = list(dict.fromkeys(unknown_people))
        errors.insert(0, AssetUploadIssue(
            sheet="전체",
            message=f"팀원 명단에 없는 사람 {len(uniq)}명: {', '.join(uniq)} — 팀원 화면에서 먼저 등록한 뒤 다시 올리세요.",
        ))

    preview = AssetUploadPreview(total=len(rows), counts=counts, rows=rows, errors=errors, warnings=warnings)
    return Plan(items=items, preview=preview)
