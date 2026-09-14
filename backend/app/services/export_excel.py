"""엑셀 export — openpyxl. TECH 09.

개별 견적서(`build_quote_xlsx`)는 원본 파일(`config.QUOTE_TEMPLATE_PATH` = 커밋된
`견적서_위탁계약용.xlsx`)을 **그대로 열어 값만 채워 넣는다** — 서식(폰트·테두리·병합·
채움색)을 코드로 재현하지 않는다. 코드로 재현했더니 폰트(Noto Sans CJK SC/맑은 고딕
vs 기본 Calibri)·테두리(굵기가 medium/thin/dotted로 셀마다 다름)가 원본과 미묘하게
달라 "육안으로 다르다"는 문제가 있었다(DECISIONS.md, tech/09-export.md) — 템플릿을
직접 여는 지금 방식은 이 문제 자체가 생기지 않는다.
공급가액/세액/합계 3행은 템플릿에 이미 있는 수식(`ROUNDDOWN(SUM(...),-5)` 등)을 그대로
살려 둬서 우리가 쓴 항목 금액에서 자동 재계산된다 — 항목이 18행(템플릿이 서식을 미리
잡아둔 행 수, `QUOTE_TEMPLATE_ITEM_ROWS`)을 넘는 경우에만 행을 늘리고 이 값들을 직접
계산해 넣는다(`_extend_item_rows`). APPROVED 견적서는 대표자명 옆에 직인(config.SEAL_PATH)
을 합성한다. 목록(`build_list_xlsx`)은 템플릿 없이 처음부터 만든다(변경 없음).
"""
from __future__ import annotations

import io
from copy import copy
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ..config import (
    COMPANY,
    COMPANY_BIZ_LINES,
    GROUPS,
    LOGO_PATH,
    QUOTE_AUTHOR_ROLE,
    QUOTE_AUTHOR_TEAM,
    QUOTE_RECIPIENT_HONORIFIC,
    QUOTE_TEMPLATE_ITEM_ROWS,
    QUOTE_TEMPLATE_PATH,
    QUOTE_TEMPLATE_SHEET,
    SEAL_MM,
    SEAL_PATH,
    STATUS_APPROVED,
    STATUS_LABELS,
)
from ..models import Quote
from .calculation import korean_amount_words
from .numbering import format_mgmt_no_display

_WON = '#,##0"원"'
_DATE_FMT = "%Y.%m.%d"

_ITEM_FIRST_ROW = 16          # 템플릿 항목 표 첫 데이터 행
_SUM_ROW_OFFSET = 34 - 16     # 합계 블록(공급가액 행)이 항목 표 첫 행에서 몇 줄 아래인지
_ITEM_MERGES = ((2, 4), (7, 8), (9, 10), (11, 12), (13, 14))  # B:D, G:H, I:J, K:L, M:N


def _fmt_date(d: date | None) -> str:
    return d.strftime(_DATE_FMT) if d else ""


def _add_logo(ws) -> int:
    """워크시트 좌상단(A1)에 회사 로고(목록 export 전용). 실패 시 조용히 건너뛴다."""
    if not LOGO_PATH.exists():
        return 1
    try:
        from openpyxl.drawing.image import Image as XLImage  # noqa: PLC0415

        img = XLImage(str(LOGO_PATH))
        img.width, img.height = 110, 48
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 36
        ws.row_dimensions[2].height = 20
        return 4
    except Exception:  # pragma: no cover - 환경 의존(Pillow 등)
        return 1


def _stamp_seal(ws, cell: str) -> None:
    """APPROVED 견적서에 직인 합성. 파일/Pillow 없으면 조용히 생략(로고와 동일 방침)."""
    if not SEAL_PATH.exists():
        return
    try:
        from openpyxl.drawing.image import Image as XLImage  # noqa: PLC0415

        px = int(SEAL_MM / 25.4 * 96)  # mm -> px @96dpi
        img = XLImage(str(SEAL_PATH))
        img.width = img.height = px
        ws.add_image(img, cell)
    except Exception:  # pragma: no cover - 환경 의존
        pass


def _extend_item_rows(ws, extra: int) -> int:
    """항목이 템플릿 행 수(18)를 넘을 때만 호출 — 합계 블록 바로 위에 행을 끼워 넣고
    마지막 항목 행(33행)의 서식(폰트·테두리·채움·병합·행높이·세액 수식)을 복제한다.
    반환값: 늘어난 뒤의 마지막 항목 행 번호.
    """
    insert_at = _ITEM_FIRST_ROW + QUOTE_TEMPLATE_ITEM_ROWS  # 34 — 합계 블록 시작 행
    template_row = insert_at - 1  # 33 — 서식 복제 원본
    ws.insert_rows(insert_at, extra)
    for offset in range(extra):
        new_row = insert_at + offset
        ws.row_dimensions[new_row].height = ws.row_dimensions[template_row].height
        for col in range(2, 15):
            src = ws.cell(row=template_row, column=col)
            dst = ws.cell(row=new_row, column=col)
            dst.font, dst.border, dst.fill = copy(src.font), copy(src.border), copy(src.fill)
            dst.alignment, dst.number_format = copy(src.alignment), src.number_format
        for c1, c2 in _ITEM_MERGES:
            ws.merge_cells(start_row=new_row, start_column=c1, end_row=new_row, end_column=c2)
        ws.cell(row=new_row, column=11, value=f"=ROUND(I{new_row}*0.1,0)")
    return template_row + extra


def build_quote_xlsx(q: Quote) -> bytes:
    wb = load_workbook(QUOTE_TEMPLATE_PATH)
    ws = wb[QUOTE_TEMPLATE_SHEET]
    ws.title = "견적서"

    group_name = GROUPS.get(q.group_code, q.group_code)

    # ---- 견적NO -------------------------------------------------------------
    mgmt_no_display = format_mgmt_no_display(q.seq_year, q.group_code, q.seq_no)
    prefix_year_group, seq_str = mgmt_no_display.rsplit(" ", 1)
    ws["M2"], ws["N2"] = prefix_year_group, seq_str

    # ---- 견적 기본정보 / 공급자(자사) — 자사정보는 config.COMPANY 가 정본이라
    #      템플릿에 이미 박혀 있는 값이라도 덮어써서 단일 출처를 지킨다 -----------
    ws["C3"] = q.issue_date.isoformat()
    ws["J3"] = COMPANY["biz_no"]
    ws["J4"] = COMPANY["name"]
    ceo_cell = "L4"
    ws[ceo_cell] = f"{COMPANY['ceo_name']} (인)"
    if q.status == STATUS_APPROVED:
        _stamp_seal(ws, ceo_cell)
    ws["J5"] = COMPANY.get("address", "")
    for i, (biz_type, biz_item) in enumerate(COMPANY_BIZ_LINES):
        row = 6 + i
        ws.cell(row=row, column=10, value=biz_type)   # J6~J8
        ws.cell(row=row, column=12, value=biz_item)    # L6~L8
    ws["J9"] = COMPANY.get("tel", "")
    ws["L9"] = COMPANY.get("fax", "")

    # ---- 수신처(고객사) -------------------------------------------------------
    ws["B4"] = q.customer_name
    ws["D4"] = q.customer_department or ""
    ws["D5"] = q.customer_contact_name or ""
    ws["F5"] = QUOTE_RECIPIENT_HONORIFIC if q.customer_contact_name else ""
    ws["B6"] = f"C.C {q.customer_cc}" if q.customer_cc else ""

    # ---- 견적 작성처 정보 -----------------------------------------------------
    ws["G10"] = f"견적 작성처 정보 : {QUOTE_AUTHOR_TEAM} {group_name} / {q.issuer_name} {QUOTE_AUTHOR_ROLE}"

    # ---- 합계금액(한글 금액 표기) — K13 은 템플릿 수식(="(₩"&TEXT(I36,...))이
    #      I36(합계) 을 참조해 그대로 재계산되므로 건드리지 않는다 -----------------
    ws["E13"] = f"일금  {korean_amount_words(q.total_with_vat)}원정"

    # ---- 항목 표: I(공급가액)만 채우면 K(세액)·I34~I36(합계 블록)·K13 이 템플릿
    #      수식으로 자동 재계산된다. 18행을 넘는 경우에만 직접 계산해 채운다 --------
    overflow = len(q.items) - QUOTE_TEMPLATE_ITEM_ROWS
    last_item_row = _extend_item_rows(ws, overflow) if overflow > 0 else _ITEM_FIRST_ROW + QUOTE_TEMPLATE_ITEM_ROWS - 1

    for idx, it in enumerate(q.items):
        row = _ITEM_FIRST_ROW + idx
        ws.cell(row=row, column=2, value=it.name)  # B
        ws.cell(row=row, column=5, value=_fmt_date(it.period_start))  # E
        ws.cell(row=row, column=6, value="~" if (it.period_start or it.period_end) else "")  # F
        ws.cell(row=row, column=7, value=_fmt_date(it.period_end))  # G
        ws.cell(row=row, column=9, value=it.line_amount).number_format = _WON  # I

    # 템플릿 자체 예시 데이터가 R16~R19 에 박혀 있어(빈 칸 아님) 실제 항목 수보다
    # 적으면 남은 행에 예시 문구가 그대로 남는다 — 명시적으로 비운다.
    # (`ws.cell(value=None)`은 openpyxl 에서 "값을 안 준 것"과 구분이 안 돼 no-op —
    # 반드시 `.value = None` 속성 대입으로 지워야 한다.)
    for row in range(_ITEM_FIRST_ROW + len(q.items), last_item_row + 1):
        for col in (2, 5, 6, 7, 9):
            ws.cell(row=row, column=col).value = None

    if overflow > 0:
        # 늘어난 행만큼 합계 수식 범위가 안 맞으므로(openpyxl 은 텍스트 수식을 자동
        # 보정하지 않는다) 세 값 직접 기록 — Quote 저장 시 계산된 정본과 동일 값.
        sum_row = last_item_row + 1
        ws.cell(row=sum_row, column=9, value=q.supply_amount).number_format = _WON
        ws.cell(row=sum_row + 1, column=9, value=q.vat_amount).number_format = _WON
        ws.cell(row=sum_row + 2, column=9, value=q.total_with_vat).number_format = _WON
        ws["K13"] = f"(₩{q.total_with_vat:,})"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_list_xlsx(quotes: list[Quote]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서 목록"

    header_row = _add_logo(ws)
    headers = [
        "관리번호", "그룹코드", "그룹명", "견적서명", "수신처", "발행일자",
        "발행담당자", "공급가액", "부가세", "부가세포함가", "부가세포함여부",
        "상태", "잠금여부", "등록시간",
    ]
    for c, hh in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=hh)
        cell.font = Font(bold=True)
    widths = [12, 8, 18, 30, 24, 12, 12, 14, 12, 14, 12, 10, 8, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w

    for offset, q in enumerate(quotes, start=1):
        row_idx = header_row + offset
        vals = [
            q.mgmt_no,
            q.group_code,
            GROUPS.get(q.group_code, ""),
            q.title,
            q.customer_name,
            q.issue_date.isoformat(),
            q.issuer_name,
            q.supply_amount,
            q.vat_amount,
            q.total_with_vat,
            "포함" if q.vat_included else "미포함",
            STATUS_LABELS.get(q.status, q.status),
            "잠금" if q.purchase_locked else "",
            q.created_at.strftime("%Y-%m-%d %H:%M") if q.created_at else "",
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=c, value=v)
            if c in (8, 9, 10):
                cell.number_format = _WON

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def list_filename() -> str:
    return f"quotes_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
