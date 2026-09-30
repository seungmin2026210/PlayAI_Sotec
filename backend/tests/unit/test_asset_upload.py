"""ASSET-1 엑셀 일괄 업로드 — 파싱·검증(순수, DB 불필요). D50~D52, P10."""
from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from app.config import ASSET_UPLOAD_MAX_ROWS
from app.errors import AppError
from app.models import Member
from app.services.asset_upload import build_plan, build_template_xlsx, columns, parse_workbook

TODAY = date(2026, 9, 30)
MEMBERS = [
    Member("Z001", "김철수", "B"),
    Member("Z002", "이영희", "A"),
    Member("Z003", "박민수", "A"),
    Member("Z004", "박민수", "C"),
    Member("Z009", "퇴사자", "A", active=False),
]


def xlsx(sheets: dict[str, list[dict]]) -> bytes:
    """{시트명: [{머리글: 값}]} → 템플릿 머리글 순서의 엑셀 바이트."""
    wb = Workbook()
    wb.remove(wb.active)
    cats = {"SW": "SW", "HW": "HW", "교육": "EDU"}
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        heads = [c.header for c in columns(cats.get(title, "SW"))]
        ws.append(heads)
        for r in rows:
            ws.append([r.get(h) for h in heads])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def plan(sheets, existing=None):
    return build_plan(parse_workbook(xlsx(sheets)), MEMBERS, existing or [], today=TODAY)


def msgs(issues):
    return [i.message for i in issues]


# --------------------------------------------------------------------------- 템플릿
def test_template_has_three_sheets_with_dropdowns():
    wb = load_workbook(io.BytesIO(build_template_xlsx()))
    assert wb.sheetnames == ["안내", "SW", "HW", "교육"]
    sw = wb["SW"]
    assert sw["A1"].value == "품명 *"
    assert "비밀번호" in [c.value for c in sw[1]]
    assert "공용 장소/용도" not in [c.value for c in wb["교육"][1]]  # 교육은 공용 배정 불가
    assert "소분류" not in [c.value for c in wb["교육"][1]]
    formulas = [dv.formula1 for dv in sw.data_validations.dataValidation]
    assert any("기술개발그룹" in f for f in formulas) and any("AI 구독" in f for f in formulas)


def test_template_roundtrips_empty_is_rejected():
    with pytest.raises(AppError, match="등록할 행이 없습니다"):
        parse_workbook(build_template_xlsx())


# --------------------------------------------------------------------------- 파일 단위 거부
def test_not_excel_rejected():
    with pytest.raises(AppError, match="읽을 수 없습니다"):
        parse_workbook(b"hello")


def test_no_category_sheet_rejected():
    wb = Workbook()
    wb.active.title = "자산목록"
    wb.active.append(["품명"])
    wb.active.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(AppError, match="시트를 찾을 수 없습니다"):
        parse_workbook(buf.getvalue())


def test_missing_required_header_rejected():
    wb = Workbook()
    wb.active.title = "SW"
    wb.active.append(["품명", "비고"])
    wb.active.append(["IntelliJ", ""])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(AppError, match="필수 열이 없습니다: 그룹"):
        parse_workbook(buf.getvalue())


def test_row_limit():
    rows = [{"품명": f"n{i}", "그룹": "A"} for i in range(ASSET_UPLOAD_MAX_ROWS + 1)]
    with pytest.raises(AppError, match=f"{ASSET_UPLOAD_MAX_ROWS}행까지"):
        parse_workbook(xlsx({"SW": rows}))


# --------------------------------------------------------------------------- 값 변환
def test_values_normalized():
    p = plan({"SW": [{
        "품명": " IntelliJ ", "소분류": "라이선스", "그룹": "기술개발그룹", "구매일": datetime(2025, 3, 10),
        "금액": "1,200,000원", "유효기간 시작": "2025.03.10", "유효기간 종료": "2026/3/9", "버전": 2024.1,
        "기존 관리번호": "OLD-7", "비고": "메모",
    }], "HW": [{"품명": "LG gram", "그룹": "B", "시리얼": 12345, "구매일": 45000}]})
    assert p.ok, p.preview.errors
    sw, hw = p.items
    f = sw.fields
    assert (f["name"], f["subcategory"], f["group_code"]) == ("IntelliJ", "LICENSE", "A")
    assert (f["purchase_date"], f["valid_from"], f["valid_to"]) == ("2025-03-10", "2025-03-10", "2026-03-09")
    assert f["price"] == 1_200_000 and f["version"] == "2024.1"
    assert f["note"] == "기존 관리번호: OLD-7\n메모"
    assert hw.fields["serial_no"] == "12345" and hw.fields["group_code"] == "B"
    assert hw.fields["purchase_date"] == "2023-03-15"  # 엑셀 날짜 일련번호
    assert p.preview.counts == {"SW": 1, "HW": 1, "EDU": 0}


def test_row_errors_collected_and_nothing_planned():
    p = plan({"SW": [
        {"품명": "ok", "그룹": "A"},
        {"그룹": "없는그룹", "금액": -5, "구매일": "2026-13-01", "소분류": "노트북"},
        {"품명": "x", "그룹": "A", "유효기간 시작": "2026-02-01", "유효기간 종료": "2026-01-01"},
        {"품명": "y" * 201, "그룹": "A"},
    ]})
    assert not p.ok
    by_row = {}
    for e in p.preview.errors:
        by_row.setdefault(e.row, []).append(e.column)
    assert 2 not in by_row
    assert set(by_row[3]) == {"품명", "그룹", "금액", "구매일", "소분류"}
    assert 4 in by_row and by_row[5] == ["품명"]
    assert p.preview.total == 4 and len(p.items) == 1  # 라우터가 errors 있으면 아무것도 저장 안 함


# --------------------------------------------------------------------------- 배정 대상
def test_member_by_id_or_unique_name():
    p = plan({"SW": [
        {"품명": "a", "그룹": "A", "사용자 사번": "Z001", "사용 시작일": "2026-01-02"},
        {"품명": "b", "그룹": "A", "사용자 이름": "이영희", "구매일": "2026-02-01"},
        {"품명": "c", "그룹": "A", "공용 장소/용도": "3층 회의실"},
        {"품명": "d", "그룹": "A"},
    ]})
    assert p.ok, p.preview.errors
    a, b, c, d = p.items
    assert (a.target.member_id, a.target.scope_group_code, a.start_date) == ("Z001", "B", "2026-01-02")
    assert (b.target.member_id, b.start_date) == ("Z002", "2026-02-01")  # 시작일 비어 → 구매일
    assert (c.target.shared_label, c.target.scope_group_code, c.start_date) == ("3층 회의실", "A", TODAY.isoformat())
    assert d.target is None and d.start_date is None
    assert p.preview.rows[1].start_auto and p.preview.rows[0].start_auto is False
    assert sum("로 채웁니다" in m for m in msgs(p.preview.warnings)) == 2


def test_member_errors():
    p = plan({"SW": [
        {"품명": "a", "그룹": "A", "사용자 사번": "Z777"},
        {"품명": "b", "그룹": "A", "사용자 이름": "홍길동"},
        {"품명": "c", "그룹": "A", "사용자 이름": "박민수"},
        {"품명": "d", "그룹": "A", "사용자 사번": "Z009"},
        {"품명": "e", "그룹": "A", "사용자 이름": "이영희", "공용 장소/용도": "회의실"},
        {"품명": "f", "그룹": "A", "사용자 사번": "Z001", "사용 시작일": "2026-10-01"},
        {"품명": "g", "그룹": "A", "사용자 사번": "Z001", "구매일": "2026-05-01", "사용 시작일": "2026-04-01"},
    ]})
    m = msgs(p.preview.errors)
    assert m[0].startswith("팀원 명단에 없는 사람 2명: Z777, 홍길동")
    assert any("같은 이름의 팀원이 여러 명" in x for x in m)
    assert any("퇴사 처리된 팀원" in x for x in m)
    assert any("중 하나만" in x for x in m)
    assert any("미래 날짜" in x for x in m)
    assert any("구매일(2026-05-01)보다 앞설 수 없습니다" in x for x in m)
    assert p.items == []


def test_name_mismatch_with_id_warns_and_uses_id():
    p = plan({"SW": [{"품명": "a", "그룹": "A", "사용자 사번": "Z001", "사용자 이름": "김철", "사용 시작일": "2026-01-01"}]})
    assert p.ok and p.items[0].target.member_name == "김철수"
    assert any("사번 기준으로 배정" in x for x in msgs(p.preview.warnings))


# --------------------------------------------------------------------------- 중복 의심 · 기타 경고
def test_duplicates_warn_only():
    existing = [
        {"asset_no": "HW-25-001", "status": "IDLE", "serial_no": "SN1"},
        {"asset_no": "HW-25-002", "status": "DELETED", "serial_no": "SN2"},  # 삭제된 건 비교 안 함
    ]
    p = plan({"HW": [
        {"품명": "a", "그룹": "A", "시리얼": "SN1"},
        {"품명": "b", "그룹": "A", "시리얼": "SN2"},
        {"품명": "c", "그룹": "A", "시리얼": "SN2"},
    ]}, existing)
    assert p.ok
    w = msgs(p.preview.warnings)
    assert any("HW-25-001" in x for x in w)
    assert sum("HW-25-002" in x for x in w) == 0
    assert any("HW 시트 3행과 같습니다" in x for x in w)


def test_unknown_column_and_orphan_start_date_warn():
    wb = Workbook()
    wb.active.title = "SW"
    wb.active.append(["품명*", "그룹 *", "담당부서", "사용 시작일"])
    wb.active.append(["a", "A", "개발", "2026-01-01"])
    buf = io.BytesIO()
    wb.save(buf)
    p = build_plan(parse_workbook(buf.getvalue()), MEMBERS, [], today=TODAY)
    assert p.ok
    w = msgs(p.preview.warnings)
    assert any("알 수 없는 열" in x for x in w) and any("시작일은 무시" in x for x in w)
    assert p.items[0].start_date is None


def test_password_kept_plain_until_save():
    p = plan({"SW": [{"품명": "a", "그룹": "A", "계정 ID": "dev01", "비밀번호": "pw!"}]})
    assert p.items[0].fields["password"] == "pw!" and p.preview.rows[0].has_password
