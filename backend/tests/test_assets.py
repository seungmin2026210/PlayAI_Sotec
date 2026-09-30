"""자산관리(ASSET-1) — specs/ASSET-1/TECH.md "테스트 계획"을 코드로 검증한다(에뮬레이터 필요)."""
from __future__ import annotations

from datetime import timedelta

import pytest
from cryptography.fernet import Fernet

from app.config import settings
from app.services.asset_dates import today_kst

TODAY = today_kst()
YY = f"{TODAY.year % 100:02d}"
D = lambda days=0: (TODAY + timedelta(days=days)).isoformat()  # noqa: E731


@pytest.fixture(autouse=True)
def _secret_key(monkeypatch):
    monkeypatch.setattr(settings, "asset_secret_key", Fernet.generate_key().decode())


# --------------------------------------------------------------------------- helpers
def _create(client, h, category="SW", quantity=1, **over):
    payload = {"category": category, "name": "GitHub Copilot Business", "group_code": "A", "quantity": quantity}
    payload.update(over)
    return client.post("/api/assets", json=payload, headers=h)


def _one(client, h, **over) -> str:
    r = _create(client, h, **over)
    assert r.status_code == 201, r.text
    return r.json()["asset_nos"][0]


def _member(client, h, no="Z001", name="김철수", group="A"):
    r = client.post("/api/members", json={"employee_no": no, "name": name, "group_code": group}, headers=h)
    assert r.status_code == 201, r.text
    return no


def _assign(client, h, asset_no, member_id=None, shared=None, start=None):
    body = {"member_id": member_id, "shared_label": shared, "start_date": start or D(-10)}
    return client.post(f"/api/assets/{asset_no}/assign", json=body, headers=h)


def _get(client, h, asset_no):
    return client.get(f"/api/assets/{asset_no}", headers=h)


def _code(r):
    return r.json()["detail"]["code"]


def _unit(client, h, category="SW", **over):
    p = client.post("/api/asset-products", json={"name": "코파일럿", "asset_category": category}, headers=h).json()
    payload = {
        "product_id": p["id"], "purchase_date": "2025-03-10", "purchased_from": "리셀러",
        "price": 1000, "unit_type": "ACCOUNT" if category != "HW" else None,
        "key_value": "sotec_dev" if category != "HW" else None, "expire_date": "2026-12-31", "group_code": "B",
    }
    payload.update(over)
    r = client.post("/api/asset-units", json=payload, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["unit_no"]


# --------------------------------------------------------------------------- 채번
def test_numbering_per_category_and_purchase_year(client, admin_headers):
    assert _one(client, admin_headers, purchase_date="2023-05-01") == "SW-23-001"
    assert _one(client, admin_headers, purchase_date="2023-07-01") == "SW-23-002"
    assert _one(client, admin_headers, category="HW", name="LG gram", purchase_date="2023-07-01") == "HW-23-001"
    assert _one(client, admin_headers) == f"SW-{YY}-001"  # 구매일 없으면 KST 올해


def test_quantity_creates_consecutive_numbers(client, admin_headers):
    r = _create(client, admin_headers, quantity=3, purchase_date="2026-01-10", price=5000)
    assert r.json()["asset_nos"] == ["SW-26-001", "SW-26-002", "SW-26-003"]
    # 직접 등록은 입력 금액을 자산마다 그대로
    assert _get(client, admin_headers, "SW-26-002").json()["price"] == 5000


@pytest.mark.parametrize("n", [0, 101])
def test_quantity_out_of_range_400(client, admin_headers, n):
    r = _create(client, admin_headers, quantity=n)
    assert r.status_code == 400 and _code(r) == "VALIDATION_ERROR"


def test_seq_exhausted_creates_nothing(client, admin_headers, db):
    db.collection("asset_seq").document("SW-26").set({"last_seq": 998})
    r = _create(client, admin_headers, quantity=2, purchase_date="2026-01-01")
    assert r.status_code == 409 and _code(r) == "SEQ_EXHAUSTED"
    assert db.collection("asset_seq").document("SW-26").get().get("last_seq") == 998
    assert not list(db.collection("assets").stream())


def test_number_does_not_change_when_purchase_date_edited(client, admin_headers):
    no = _one(client, admin_headers, purchase_date="2024-01-01")
    r = client.patch(f"/api/assets/{no}", json={"purchase_date": "2026-01-01"}, headers=admin_headers)
    assert r.status_code == 200
    assert r.json()["asset_no"] == "SW-24-001" and r.json()["purchase_date"] == "2026-01-01"


# --------------------------------------------------------------------------- 수정 규칙
def test_category_cannot_change(client, admin_headers):
    no = _one(client, admin_headers)
    r = client.patch(f"/api/assets/{no}", json={"category": "HW"}, headers=admin_headers)
    assert r.status_code == 400


def test_password_blank_keeps_and_delete_removes(client, admin_headers):
    no = _one(client, admin_headers, account_id="dev01", password="s3cret")
    assert _get(client, admin_headers, no).json()["has_password"] is True
    client.patch(f"/api/assets/{no}", json={"password": "", "note": "메모"}, headers=admin_headers)
    assert _get(client, admin_headers, no).json()["has_password"] is True
    r = client.post(f"/api/assets/{no}/password/reveal", json={"action": "REVEAL"}, headers=admin_headers)
    assert r.json()["password"] == "s3cret"
    client.delete(f"/api/assets/{no}/password", headers=admin_headers)
    assert _get(client, admin_headers, no).json()["has_password"] is False


# --------------------------------------------------------------------------- 비밀번호
def test_password_never_serialized(client, admin_headers, manager_headers, db):
    no = _one(client, admin_headers, account_id="dev01", password="s3cret", license_key="ABCD-EFGH")
    enc = db.collection("assets").document(no).get().get("password_enc")
    assert enc and "s3cret" not in enc
    for r in (
        _get(client, admin_headers, no),
        client.get("/api/assets?category=SW", headers=admin_headers),
        client.get("/api/assets/export.xlsx?category=SW", headers=admin_headers),
        _get(client, manager_headers, no),
    ):
        assert r.status_code == 200
        body = r.content
        assert b"s3cret" not in body and enc.encode() not in body and b"password_enc" not in body


def test_reveal_logs_and_permissions(client, admin_headers, manager_headers, db):
    no = _one(client, admin_headers, account_id="dev01", password="pw")
    r = client.post(f"/api/assets/{no}/password/reveal", json={"action": "COPY"}, headers=admin_headers)
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    logs = [d.to_dict() for d in db.collection("password_reveal_logs").stream()]
    assert len(logs) == 1 and logs[0]["username"] == "Admin" and logs[0]["action"] == "COPY"
    r = client.post(f"/api/assets/{no}/password/reveal", json={"action": "REVEAL"}, headers=manager_headers)
    assert r.status_code == 403


def test_secret_key_missing(client, admin_headers, monkeypatch, db):
    monkeypatch.setattr(settings, "asset_secret_key", None)
    r = _create(client, admin_headers, account_id="dev01", password="pw")
    assert r.status_code == 501 and _code(r) == "SECRET_KEY_MISSING"
    assert not list(db.collection("assets").stream())  # 통째로 실패(D43)
    assert _create(client, admin_headers, account_id="dev01").status_code == 201  # 비밀번호 없으면 정상


# --------------------------------------------------------------------------- 배정 · 회수 · 이관
def test_assign_return_sync_scope(client, admin_headers, db):
    no = _one(client, admin_headers, group_code="A")
    _member(client, admin_headers, "Z001", "김철수", "B")
    r = _assign(client, admin_headers, no, member_id="Z001")
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["status"] == "IN_USE" and a["current_member_name"] == "김철수" and a["scope_group_code"] == "B"
    assert len(a["assignments"]) == 1

    assert _code(_assign(client, admin_headers, no, member_id="Z001")) == "ASSET_NOT_IDLE"

    r = client.post(f"/api/assets/{no}/return", json={"end_date": D(-1)}, headers=admin_headers)
    a = r.json()
    assert a["status"] == "IDLE" and a["current_member_id"] is None and a["scope_group_code"] == "A"
    assert a["assignments"][0]["end_date"] == D(-1)


def test_shared_assignment(client, admin_headers):
    hw = _one(client, admin_headers, category="HW", name="공유기", group_code="C")
    a = _assign(client, admin_headers, hw, shared="3층 회의실").json()
    assert a["status"] == "IN_USE" and a["current_shared_label"] == "3층 회의실" and a["scope_group_code"] == "C"
    edu = _one(client, admin_headers, category="EDU", name="스프링")
    assert _assign(client, admin_headers, edu, shared="회의실").status_code == 400
    listed = client.get("/api/assets?category=HW&member_id=SHARED", headers=admin_headers).json()
    assert [i["asset_no"] for i in listed["items"]] == [hw]


def test_member_or_shared_exactly_one(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers)
    assert _assign(client, admin_headers, no).status_code == 400
    assert _assign(client, admin_headers, no, member_id="Z001", shared="x").status_code == 400


def test_inactive_member_cannot_be_assigned(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers)
    client.patch("/api/members/Z001", json={"active": False}, headers=admin_headers)
    r = _assign(client, admin_headers, no, member_id="Z001")
    assert r.status_code == 409 and _code(r) == "MEMBER_INACTIVE"


def test_transfer_is_atomic_and_cancel_restores_previous(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers, "Z001", "김철수", "A")
    _member(client, admin_headers, "Z002", "이영희", "B")
    _assign(client, admin_headers, no, member_id="Z001", start=D(-10))

    r = client.post(f"/api/assets/{no}/transfer", json={"member_id": "Z002", "date": D(-3)}, headers=admin_headers)
    a = r.json()
    assert r.status_code == 200, r.text
    assert a["current_member_id"] == "Z002" and a["scope_group_code"] == "B" and a["current_start_date"] == D(-3)
    ends = {x["member_id"]: x["end_date"] for x in a["assignments"]}
    assert ends == {"Z001": D(-3), "Z002": None}

    a = client.post(f"/api/assets/{no}/assignment/cancel", headers=admin_headers).json()
    assert a["current_member_id"] == "Z001" and a["scope_group_code"] == "A" and a["current_start_date"] == D(-10)
    assert [(x["member_id"], x["end_date"]) for x in a["assignments"]] == [("Z001", None)]

    a = client.post(f"/api/assets/{no}/assignment/cancel", headers=admin_headers).json()
    assert a["status"] == "IDLE" and a["assignments"] == []
    r = client.post(f"/api/assets/{no}/assignment/cancel", headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ASSET_NOT_IN_USE"


def test_transfer_date_before_start_400(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers, "Z001")
    _member(client, admin_headers, "Z002", "이영희")
    _assign(client, admin_headers, no, member_id="Z001", start=D(-5))
    r = client.post(f"/api/assets/{no}/transfer", json={"member_id": "Z002", "date": D(-6)}, headers=admin_headers)
    assert r.status_code == 400


@pytest.mark.parametrize("action", ["assign", "return", "transfer"])
def test_future_date_rejected(client, admin_headers, action):
    no = _one(client, admin_headers)
    _member(client, admin_headers, "Z001")
    _member(client, admin_headers, "Z002", "이영희")
    if action == "assign":
        r = _assign(client, admin_headers, no, member_id="Z001", start=D(1))
    else:
        _assign(client, admin_headers, no, member_id="Z001", start=D(-5))
        if action == "return":
            r = client.post(f"/api/assets/{no}/return", json={"end_date": D(1)}, headers=admin_headers)
        else:
            r = client.post(f"/api/assets/{no}/transfer", json={"member_id": "Z002", "date": D(1)}, headers=admin_headers)
    assert r.status_code == 400 and _code(r) == "VALIDATION_ERROR"


def test_period_overlap_and_return_before_start(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers, "Z001")
    _assign(client, admin_headers, no, member_id="Z001", start=D(-10))
    assert client.post(f"/api/assets/{no}/return", json={"end_date": D(-11)}, headers=admin_headers).status_code == 400
    client.post(f"/api/assets/{no}/return", json={"end_date": D(-5)}, headers=admin_headers)
    # 앞 사용자 종료일 이전 날짜로 새 배정 불가
    assert _assign(client, admin_headers, no, member_id="Z001", start=D(-7)).status_code == 400
    assert _assign(client, admin_headers, no, member_id="Z001", start=D(-5)).status_code == 200


def test_edit_assignment(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers, "Z001")
    a = _assign(client, admin_headers, no, member_id="Z001", start=D(-10)).json()
    aid = a["assignments"][0]["id"]
    # 활성 이력의 종료일은 수정으로 못 채움
    assert client.patch(f"/api/asset-assignments/{aid}", json={"end_date": D(-1)}, headers=admin_headers).status_code == 400
    assert client.patch(f"/api/asset-assignments/{aid}", json={"start_date": D(1)}, headers=admin_headers).status_code == 400
    r = client.patch(f"/api/asset-assignments/{aid}", json={"start_date": D(-8), "note": "수기 정정"}, headers=admin_headers)
    a = r.json()
    assert a["current_start_date"] == D(-8)
    row = a["assignments"][0]
    assert row["start_date"] == D(-8) and row["note"] == "수기 정정" and row["updated_by"] == "Admin"
    assert row["member_id"] == "Z001"  # 사람은 그대로


# --------------------------------------------------------------------------- 폐기
def test_dispose_rules(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers)
    _assign(client, admin_headers, no, member_id="Z001")
    r = client.post(f"/api/assets/{no}/dispose", headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ASSET_IN_USE"
    client.post(f"/api/assets/{no}/return", json={"end_date": D()}, headers=admin_headers)
    a = client.post(f"/api/assets/{no}/dispose", headers=admin_headers).json()
    assert a["status"] == "DISPOSED" and a["read_only"] is True and len(a["assignments"]) == 1

    assert _code(client.patch(f"/api/assets/{no}", json={"note": "x"}, headers=admin_headers)) == "ASSET_DISPOSED"
    assert _code(_assign(client, admin_headers, no, member_id="Z001")) == "ASSET_DISPOSED"
    # 기본 목록에서 숨김, 상태 필터로는 보임
    assert client.get("/api/assets?category=SW", headers=admin_headers).json()["total"] == 0
    assert client.get("/api/assets?category=SW&status=DISPOSED", headers=admin_headers).json()["total"] == 1


# --------------------------------------------------------------------------- 삭제(D47)
def _delete(client, h, asset_no, reason="중복 등록"):
    return client.post(f"/api/assets/{asset_no}/delete", json={"reason": reason}, headers=h)


def test_delete_records_who_when_why_and_hides(client, admin_headers, manager_headers):
    keep = _one(client, admin_headers, valid_to=D(5))
    no = _one(client, admin_headers, valid_to=D(5))
    assert _delete(client, admin_headers, no, reason="  ").status_code == 422  # 사유 필수
    a = _delete(client, admin_headers, no).json()
    assert a["status"] == "DELETED" and a["read_only"] is True
    assert a["deleted_by"] == "Admin" and a["deleted_reason"] == "중복 등록" and a["deleted_at"]

    assert _code(client.patch(f"/api/assets/{no}", json={"note": "x"}, headers=admin_headers)) == "ASSET_DELETED"
    assert _code(client.post(f"/api/assets/{no}/dispose", headers=admin_headers)) == "ASSET_DELETED"
    assert _code(_delete(client, admin_headers, no)) == "ASSET_DELETED"
    # 기본 목록·대시보드에서 숨김, "삭제됨" 필터는 전체관리자만
    assert [i["asset_no"] for i in client.get("/api/assets?category=SW", headers=admin_headers).json()["items"]] == [keep]
    assert client.get("/api/assets?category=SW&status=DELETED", headers=admin_headers).json()["total"] == 1
    assert client.get("/api/assets?category=SW&status=DELETED", headers=manager_headers).json()["total"] == 0
    assert _get(client, manager_headers, no).status_code == 404
    kpi = client.get(f"/api/dashboard/renewals?from={D(-30)}&to={D(30)}", headers=admin_headers).json()["kpi"]
    assert kpi["d30"] == 1
    # 번호는 결번 — 새 등록은 다음 번호
    assert _one(client, admin_headers) == f"SW-{YY}-003"


def test_delete_rules(client, admin_headers, manager_headers):
    no = _one(client, admin_headers, valid_to="2026-12-31")
    assert _delete(client, manager_headers, no).status_code == 403
    # 배정 이력(회수 후에도)
    _member(client, admin_headers)
    _assign(client, admin_headers, no, member_id="Z001", start=D(-3))
    assert _code(_delete(client, admin_headers, no)) == "ASSET_HAS_HISTORY"
    client.post(f"/api/assets/{no}/return", json={"end_date": D()}, headers=admin_headers)
    assert _code(_delete(client, admin_headers, no)) == "ASSET_HAS_HISTORY"
    # 갱신 기록
    renewed = _one(client, admin_headers, valid_to="2026-12-31")
    client.post("/api/assets/renew", json={"asset_nos": [renewed], "new_valid_to": "2027-12-31"}, headers=admin_headers)
    assert _code(_delete(client, admin_headers, renewed)) == "ASSET_HAS_HISTORY"
    # 폐기된 자산
    disposed = _one(client, admin_headers)
    client.post(f"/api/assets/{disposed}/dispose", headers=admin_headers)
    assert _code(_delete(client, admin_headers, disposed)) == "ASSET_DISPOSED"
    # 구매에서 가져온 자산 → 가져오기 취소로
    unit_no = _unit(client, admin_headers)
    imported = client.post("/api/assets/import", json={"unit_no": unit_no}, headers=admin_headers).json()["asset_nos"][0]
    assert _code(_delete(client, admin_headers, imported)) == "ASSET_IMPORTED"


# --------------------------------------------------------------------------- 가져오기
def test_import_maps_values_and_splits_price(client, admin_headers, db):
    unit_no = _unit(client, admin_headers, price=1000)
    r = client.post("/api/assets/import", json={"unit_no": unit_no, "quantity": 3, "subcategory": "AI_SUB"}, headers=admin_headers)
    assert r.status_code == 201, r.text
    nos = r.json()["asset_nos"]
    assert nos == ["SW-25-001", "SW-25-002", "SW-25-003"]
    items = [_get(client, admin_headers, n).json() for n in nos]
    assert [i["price"] for i in items] == [334, 333, 333]
    first = items[0]
    assert first["name"] == "코파일럿" and first["account_id"] == "sotec_dev" and first["valid_to"] == "2026-12-31"
    assert first["purchased_from"] == "리셀러" and first["group_code"] == "B" and first["source_unit_no"] == unit_no

    unit = client.get(f"/api/asset-units/{unit_no}", headers=admin_headers).json()
    assert unit["asset_link_kind"] == "IMPORTED" and unit["asset_nos"] == nos and unit["asset_link_label"] == "자산 등록됨"
    r = client.post("/api/assets/import", json={"unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ALREADY_IMPORTED"
    assert client.get("/api/assets/importable?category=SW", headers=admin_headers).json() == []


def test_import_retired_unit_409(client, admin_headers):
    unit_no = _unit(client, admin_headers)
    client.post(f"/api/asset-units/{unit_no}/retire", headers=admin_headers)
    r = client.post("/api/assets/import", json={"unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ALREADY_RETIRED"


def test_cancel_import(client, admin_headers):
    unit_no = _unit(client, admin_headers)
    nos = client.post("/api/assets/import", json={"unit_no": unit_no, "quantity": 2}, headers=admin_headers).json()["asset_nos"]
    _member(client, admin_headers)
    _assign(client, admin_headers, nos[1], member_id="Z001")
    r = client.post("/api/assets/import/cancel", json={"unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ASSET_HAS_HISTORY"

    client.post(f"/api/assets/{nos[1]}/assignment/cancel", headers=admin_headers)
    r = client.post("/api/assets/import/cancel", json={"unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 200, r.text
    a = _get(client, admin_headers, nos[0]).json()
    assert a["status"] == "DISPOSED" and a["disposed_reason"] == "IMPORT_CANCELLED"
    assert client.get(f"/api/asset-units/{unit_no}", headers=admin_headers).json()["asset_link_kind"] is None
    # 다시 가져오면 새 번호(결번 유지)
    again = client.post("/api/assets/import", json={"unit_no": unit_no}, headers=admin_headers).json()["asset_nos"]
    assert again == ["SW-25-003"]


def test_cancel_import_rejects_history_after_return(client, admin_headers):
    unit_no = _unit(client, admin_headers)
    no = client.post("/api/assets/import", json={"unit_no": unit_no}, headers=admin_headers).json()["asset_nos"][0]
    _member(client, admin_headers)
    _assign(client, admin_headers, no, member_id="Z001", start=D(-3))
    client.post(f"/api/assets/{no}/return", json={"end_date": D()}, headers=admin_headers)
    assert _code(client.post("/api/assets/import/cancel", json={"unit_no": unit_no}, headers=admin_headers)) == "ASSET_HAS_HISTORY"


# --------------------------------------------------------------------------- 구매 폐기 연쇄(C4)
def test_retire_unit_disposes_imported_assets(client, admin_headers):
    unit_no = _unit(client, admin_headers)
    nos = client.post("/api/assets/import", json={"unit_no": unit_no, "quantity": 2}, headers=admin_headers).json()["asset_nos"]
    client.post(f"/api/assets/{nos[0]}/dispose", headers=admin_headers)  # 이미 폐기 — 사유 유지
    r = client.post(f"/api/asset-units/{unit_no}/retire", headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "EXPIRED"
    a0, a1 = (_get(client, admin_headers, n).json() for n in nos)
    assert a0["status"] == "DISPOSED" and a0["disposed_reason"] == "DISPOSED"
    assert a1["status"] == "DISPOSED" and a1["disposed_reason"] == "PURCHASE_RETIRED"
    assert a1["disposed_reason_label"] == "구매 폐기"
    # 폐기된 구매 기록은 가져오기 취소 불가
    r = client.post("/api/assets/import/cancel", json={"unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ALREADY_RETIRED"


def test_retire_unit_blocked_by_in_use_asset(client, admin_headers):
    unit_no = _unit(client, admin_headers)
    nos = client.post("/api/assets/import", json={"unit_no": unit_no, "quantity": 2}, headers=admin_headers).json()["asset_nos"]
    _member(client, admin_headers)
    _assign(client, admin_headers, nos[1], member_id="Z001")
    r = client.post(f"/api/asset-units/{unit_no}/retire", headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "ASSET_IN_USE" and nos[1] in r.json()["detail"]["message"]
    # 전부 아니면 전무 — 아무것도 안 바뀜
    assert client.get(f"/api/asset-units/{unit_no}", headers=admin_headers).json()["retired_at"] is None
    assert _get(client, admin_headers, nos[0]).json()["status"] == "IDLE"


def test_retire_unit_keeps_renewed_assets(client, admin_headers):
    nos = _create(client, admin_headers, quantity=2, valid_from="2026-01-01", valid_to="2026-12-31").json()["asset_nos"]
    unit_no = _unit(client, admin_headers)
    client.post("/api/assets/renew", json={"asset_nos": nos, "new_valid_to": "2027-12-31", "unit_no": unit_no}, headers=admin_headers)
    assert client.post(f"/api/asset-units/{unit_no}/retire", headers=admin_headers).status_code == 200
    assert all(_get(client, admin_headers, n).json()["status"] == "IDLE" for n in nos)


# --------------------------------------------------------------------------- 갱신
def test_renew_bulk_with_unit(client, admin_headers):
    nos = _create(client, admin_headers, quantity=2, valid_from="2026-01-01", valid_to="2026-12-31").json()["asset_nos"]
    unit_no = _unit(client, admin_headers)
    r = client.post("/api/assets/renew", json={"asset_nos": nos, "new_valid_to": "2027-12-31", "unit_no": unit_no}, headers=admin_headers)
    assert r.status_code == 200, r.text
    a = _get(client, admin_headers, nos[0]).json()
    assert a["valid_from"] == "2027-01-01" and a["valid_to"] == "2027-12-31"
    assert a["renewals"][0]["prev_valid_to"] == "2026-12-31" and a["renewals"][0]["count"] == 2
    unit = client.get(f"/api/asset-units/{unit_no}", headers=admin_headers).json()
    assert unit["asset_link_kind"] == "RENEWED" and unit["asset_link_label"] == "갱신에 사용됨"
    assert client.get("/api/assets/importable?category=SW", headers=admin_headers).json() == []


def test_renew_validation(client, admin_headers):
    a = _one(client, admin_headers, valid_to="2026-12-31")
    b = _one(client, admin_headers, valid_to="2026-11-30")
    c = _one(client, admin_headers)
    renew = lambda nos, to: client.post("/api/assets/renew", json={"asset_nos": nos, "new_valid_to": to}, headers=admin_headers)  # noqa: E731
    assert renew([a, b], "2027-12-31").status_code == 400  # 종료일 다른 묶음
    assert renew([c], "2027-12-31").status_code == 400  # 종료일 없음
    assert renew([a], "2026-12-31").status_code == 400  # 새 종료일 ≤ 기존
    client.post(f"/api/assets/{a}/dispose", headers=admin_headers)
    assert _code(renew([a], "2027-12-31")) == "ASSET_DISPOSED"


# --------------------------------------------------------------------------- 팀원
def test_member_rename_and_regroup_propagate(client, admin_headers):
    no = _one(client, admin_headers, group_code="A")
    _member(client, admin_headers, "Z001", "김철수", "A")
    _assign(client, admin_headers, no, member_id="Z001")
    r = client.patch("/api/members/Z001", json={"name": "김철민", "group_code": "C"}, headers=admin_headers)
    assert r.status_code == 200
    a = _get(client, admin_headers, no).json()
    assert a["current_member_name"] == "김철민" and a["scope_group_code"] == "C"
    assert a["assignments"][0]["member_name"] == "김철수"  # 이력은 스냅샷


def test_member_retire_warns(client, admin_headers):
    no = _one(client, admin_headers)
    _member(client, admin_headers)
    _assign(client, admin_headers, no, member_id="Z001")
    r = client.patch("/api/members/Z001", json={"active": False}, headers=admin_headers).json()
    assert r["active"] is False and "1건" in r["warning"]
    assert _get(client, admin_headers, no).json()["status"] == "IN_USE"  # 자동 회수 없음(P2)


def test_member_duplicate_409(client, admin_headers):
    _member(client, admin_headers)
    r = client.post("/api/members", json={"employee_no": "Z001", "name": "x", "group_code": "A"}, headers=admin_headers)
    assert r.status_code == 409 and _code(r) == "MEMBER_EXISTS"


def test_member_list_counts_and_detail(client, admin_headers):
    _member(client, admin_headers, "Z001", "김철수", "A")
    sw = _one(client, admin_headers)
    hw = _one(client, admin_headers, category="HW", name="gram")
    _assign(client, admin_headers, sw, member_id="Z001", start=D(-10))
    _assign(client, admin_headers, hw, member_id="Z001", start=D(-10))
    client.post(f"/api/assets/{hw}/return", json={"end_date": D(-2)}, headers=admin_headers)
    m = client.get("/api/members", headers=admin_headers).json()["items"][0]
    assert m["counts"] == {"SW": 1, "HW": 0, "EDU": 0}
    d = client.get("/api/members/Z001", headers=admin_headers).json()
    assert [(x["asset_no"], x["end_date"]) for x in d["assignments"]] == [(sw, None), (hw, D(-2))]


# --------------------------------------------------------------------------- 권한 · 스코프
def test_group_manager_cannot_write(client, admin_headers, manager_headers):
    no = _one(client, admin_headers)
    assert _create(client, manager_headers).status_code == 403
    assert client.patch(f"/api/assets/{no}", json={"note": "x"}, headers=manager_headers).status_code == 403
    assert client.post(f"/api/assets/{no}/dispose", headers=manager_headers).status_code == 403
    assert client.post(f"/api/assets/{no}/delete", json={"reason": "x"}, headers=manager_headers).status_code == 403
    assert client.post("/api/members", json={"employee_no": "Z9", "name": "x", "group_code": "A"}, headers=manager_headers).status_code == 403
    assert client.get("/api/assets/importable", headers=manager_headers).status_code == 403


def test_scope_follows_current_user_group(client, admin_headers, manager_headers):
    """manager(test1) = A 그룹. B 그룹 등록 자산도 A 팀원이 쓰면 보이고, 회수하면 안 보인다."""
    no = _one(client, admin_headers, group_code="B")
    assert _get(client, manager_headers, no).status_code == 404
    _member(client, admin_headers, "Z001", "김철수", "A")
    _assign(client, admin_headers, no, member_id="Z001")
    assert _get(client, manager_headers, no).status_code == 200
    assert client.get("/api/assets?category=SW", headers=manager_headers).json()["total"] == 1
    client.post(f"/api/assets/{no}/return", json={"end_date": D()}, headers=admin_headers)
    assert _get(client, manager_headers, no).status_code == 404
    assert client.get("/api/assets?category=SW", headers=manager_headers).json()["total"] == 0


def test_manager_members_own_group_and_linkable(client, admin_headers, manager_headers):
    _member(client, admin_headers, "Z001", "김철수", "A")
    _member(client, admin_headers, "Z002", "이영희", "B")
    assert [m["employee_no"] for m in client.get("/api/members", headers=manager_headers).json()["items"]] == ["Z001"]
    assert client.get("/api/members/Z002", headers=manager_headers).status_code == 404

    no = _one(client, admin_headers, group_code="A")
    _assign(client, admin_headers, no, member_id="Z001", start=D(-10))
    client.post(f"/api/assets/{no}/transfer", json={"member_id": "Z002", "date": D(-2)}, headers=admin_headers)
    rows = client.get("/api/members/Z001", headers=manager_headers).json()["assignments"]
    assert rows[0]["asset_no"] == no and rows[0]["linkable"] is False  # 지금 B 그룹 스코프(D37)
    rows = client.get("/api/members/Z001", headers=admin_headers).json()["assignments"]
    assert rows[0]["linkable"] is True


def test_license_key_masked_for_manager(client, admin_headers, manager_headers):
    no = _one(client, admin_headers, license_key="ABCD-EFGH-IJKL")
    assert _get(client, admin_headers, no).json()["license_key"] == "ABCD-EFGH-IJKL"
    assert _get(client, manager_headers, no).json()["license_key"] == "ABCD-****"
    listed = client.get("/api/assets?category=SW", headers=manager_headers).json()["items"][0]
    assert listed["license_key"] == "ABCD-****"
    # 키 검색은 전체관리자만
    assert client.get("/api/assets?category=SW&q=EFGH", headers=admin_headers).json()["total"] == 1
    assert client.get("/api/assets?category=SW&q=EFGH", headers=manager_headers).json()["total"] == 0


def test_license_key_masked_in_excel(client, admin_headers, manager_headers):
    import io

    from openpyxl import load_workbook

    _one(client, admin_headers, license_key="ABCD-EFGH-IJKL")
    ws = load_workbook(io.BytesIO(client.get("/api/assets/export.xlsx?category=SW", headers=manager_headers).content)).active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "ABCD-****" in values and "ABCD-EFGH-IJKL" not in values


# --------------------------------------------------------------------------- 목록 필터
def test_list_filters(client, admin_headers):
    _one(client, admin_headers, name="A제품", valid_to=D(-1), purchase_date="2024-02-02")
    _one(client, admin_headers, name="B제품", valid_to=D(10), subcategory="AI_SUB")
    _one(client, admin_headers, name="C제품", valid_to=D(200), note="특이사항")
    get = lambda qs: [i["name"] for i in client.get(f"/api/assets?category=SW&{qs}", headers=admin_headers).json()["items"]]  # noqa: E731
    assert get("expiry=expired") == ["A제품"]
    assert get("expiry=expiring") == ["B제품"]
    assert get("expiry=ok") == ["C제품"]
    assert get("year=2024") == ["A제품"]
    assert get("subcategory=AI_SUB") == ["B제품"]
    assert get("q=특이") == ["C제품"]
    assert get(f"name=B제품&valid_to={D(10)}") == ["B제품"]
    item = client.get("/api/assets?category=SW&expiry=expired", headers=admin_headers).json()["items"][0]
    assert item["expiry_badge"] == "EXPIRED"


# --------------------------------------------------------------------------- 대시보드
def test_dashboard_renewals_grouped_scoped(client, admin_headers, manager_headers):
    _create(client, admin_headers, quantity=3, name="Copilot", valid_to=D(5), group_code="A")
    _one(client, admin_headers, name="Copilot", valid_to=D(5), group_code="B")
    _one(client, admin_headers, name="MATLAB", valid_to=D(-3), group_code="A")
    disposed = _one(client, admin_headers, name="Old", valid_to=D(3), group_code="A")
    client.post(f"/api/assets/{disposed}/dispose", headers=admin_headers)

    r = client.get(f"/api/dashboard/renewals?from={D(-30)}&to={D(30)}", headers=admin_headers).json()
    assert [(i["name"], i["count"]) for i in r["items"]] == [("MATLAB", 1), ("Copilot", 4)]
    assert r["kpi"] == {"expired": 1, "d30": 4, "d90": 4}

    r = client.get(f"/api/dashboard/renewals?from={D(-30)}&to={D(30)}", headers=manager_headers).json()
    assert [(i["name"], i["count"]) for i in r["items"]] == [("MATLAB", 1), ("Copilot", 3)]
    assert r["kpi"]["d30"] == 3


def test_meta_has_asset_constants(client, admin_headers):
    m = client.get("/api/meta", headers=admin_headers).json()
    assert {s["code"] for s in m["asset_subcategories"]["SW"]} == {"LICENSE", "AI_SUB", "ETC"}
    assert m["asset_subcategories"]["EDU"] == []
    assert [s["code"] for s in m["asset_statuses"]] == ["IDLE", "IN_USE", "DISPOSED", "DELETED"]
