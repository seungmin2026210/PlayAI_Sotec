"""ASSET-1 엑셀 일괄 업로드 API — D50~D52, P10(에뮬레이터 필요)."""
from __future__ import annotations

import io

import pytest
from cryptography.fernet import Fernet
from openpyxl import load_workbook

from app.config import ASSET_UPLOAD_MAX_ROWS, settings
from app.services.asset_dates import today_kst
from tests.unit.test_asset_upload import xlsx

TODAY = today_kst().isoformat()


@pytest.fixture(autouse=True)
def _secret_key(monkeypatch):
    monkeypatch.setattr(settings, "asset_secret_key", Fernet.generate_key().decode())


def _member(client, h, no, name, group="A"):
    r = client.post("/api/members", json={"employee_no": no, "name": name, "group_code": group}, headers=h)
    assert r.status_code == 201, r.text


def _files(content: bytes):
    return {"file": ("assets.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def _preview(client, h, content):
    return client.post("/api/assets/upload/preview", files=_files(content), headers=h)


def _upload(client, h, content):
    return client.post("/api/assets/upload", files=_files(content), headers=h)


SAMPLE = {
    "SW": [
        {"품명": "IntelliJ", "소분류": "라이선스", "그룹": "기술개발그룹", "구매일": "2025-03-10", "금액": 500000,
         "계정 ID": "dev01", "비밀번호": "pw!", "사용자 이름": "김철수", "기존 관리번호": "OLD-1"},
        {"품명": "Copilot", "그룹": "A", "구매일": "2025-04-01"},
    ],
    "HW": [{"품명": "LG gram", "소분류": "노트북", "그룹": "A", "구매일": "2025-04-15", "시리얼": "SN1", "공용 장소/용도": "3층 회의실",
            "사용 시작일": "2025-05-01"}],
}


def test_template_download_super_admin_only(client, admin_headers, manager_headers):
    r = client.get("/api/assets/upload/template.xlsx", headers=admin_headers)
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["안내", "SW", "HW", "교육"]
    assert client.get("/api/assets/upload/template.xlsx", headers=manager_headers).status_code == 403


def test_group_manager_cannot_upload(client, manager_headers):
    content = xlsx({"SW": [{"품명": "a", "그룹": "A"}]})
    assert _preview(client, manager_headers, content).status_code == 403
    assert _upload(client, manager_headers, content).status_code == 403


def test_preview_saves_nothing_then_upload_creates_assets_and_assignments(client, admin_headers, db):
    _member(client, admin_headers, "Z001", "김철수", "B")
    content = xlsx(SAMPLE)

    r = _preview(client, admin_headers, content)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["errors"] == [] and p["counts"] == {"SW": 2, "HW": 1, "EDU": 0}
    assert p["rows"][0]["user_label"] == "김철수(Z001)" and p["rows"][0]["start_date"] == "2025-03-10"
    assert list(db.collection("assets").stream()) == []

    r = _upload(client, admin_headers, content)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["asset_nos"] == ["SW-25-001", "SW-25-002", "HW-25-001"]

    sw1 = client.get("/api/assets/SW-25-001", headers=admin_headers).json()
    assert sw1["status"] == "IN_USE" and sw1["current_member_id"] == "Z001"
    assert sw1["scope_group_code"] == "B" and sw1["current_start_date"] == "2025-03-10"
    assert sw1["has_password"] is True and sw1["note"] == "기존 관리번호: OLD-1"
    assert sw1["source_unit_no"] is None  # 직접 등록 취급(D47 삭제 가능)
    assert len(sw1["assignments"]) == 1
    hw = client.get("/api/assets/HW-25-001", headers=admin_headers).json()
    assert hw["current_shared_label"] == "3층 회의실" and hw["current_start_date"] == "2025-05-01"
    assert client.get("/api/assets/SW-25-002", headers=admin_headers).json()["status"] == "IDLE"

    docs = [d.to_dict() for d in db.collection("assets").stream()]
    assert {d["upload_batch_id"] for d in docs} == {body["batch_id"]}
    # 비밀번호는 암호문으로만
    assert db.collection("assets").document("SW-25-001").get().get("password_enc") != "pw!"

    # 업로드한 자산도 이력이 없으면 삭제 가능(D47)
    r = client.post("/api/assets/SW-25-002/delete", json={"reason": "잘못 올림"}, headers=admin_headers)
    assert r.status_code == 200, r.text


def test_upload_with_errors_rejects_everything(client, admin_headers, db):
    content = xlsx({"SW": [{"품명": "ok", "그룹": "A"}, {"품명": "bad", "그룹": "A", "사용자 이름": "홍길동"}]})
    p = _preview(client, admin_headers, content).json()
    assert "팀원 명단에 없는 사람 1명: 홍길동" in p["errors"][0]["message"]
    r = _upload(client, admin_headers, content)
    assert r.status_code == 400 and r.json()["detail"]["code"] == "ASSET_UPLOAD_INVALID"
    assert list(db.collection("assets").stream()) == []
    assert list(db.collection("asset_seq").stream()) == []


def test_file_level_error_is_400(client, admin_headers):
    r = _preview(client, admin_headers, b"not excel")
    assert r.status_code == 400 and r.json()["detail"]["code"] == "VALIDATION_ERROR"


def test_password_without_key_is_501_before_saving(client, admin_headers, monkeypatch, db):
    monkeypatch.setattr(settings, "asset_secret_key", "")
    content = xlsx({"SW": [{"품명": "a", "그룹": "A", "비밀번호": "pw"}]})
    r = _preview(client, admin_headers, content)
    assert r.status_code == 501 and r.json()["detail"]["code"] == "SECRET_KEY_MISSING"
    assert _upload(client, admin_headers, content).status_code == 501
    assert list(db.collection("assets").stream()) == []
    # 비밀번호 없으면 키 없이도 정상
    assert _upload(client, admin_headers, xlsx({"SW": [{"품명": "a", "그룹": "A"}]})).status_code == 201


def test_numbering_continues_after_existing_and_warns_duplicates(client, admin_headers):
    r = client.post("/api/assets", json={"category": "HW", "name": "gram", "group_code": "A",
                                         "purchase_date": "2025-01-01", "serial_no": "SN1"}, headers=admin_headers)
    assert r.status_code == 201
    content = xlsx({"HW": [{"품명": "gram2", "그룹": "A", "구매일": "2025-02-01", "시리얼": "SN1"}]})
    p = _preview(client, admin_headers, content).json()
    assert any("HW-25-001" in w["message"] for w in p["warnings"])
    assert _upload(client, admin_headers, content).json()["asset_nos"] == ["HW-25-002"]


def test_member_deactivated_between_preview_and_upload(client, admin_headers, db):
    _member(client, admin_headers, "Z001", "김철수")
    content = xlsx({"SW": [{"품명": "a", "그룹": "A", "사용자 사번": "Z001", "사용 시작일": TODAY}]})
    assert _preview(client, admin_headers, content).json()["errors"] == []
    client.patch("/api/members/Z001", json={"active": False}, headers=admin_headers)
    r = _upload(client, admin_headers, content)
    assert r.status_code == 400  # 재검증에서 걸림
    assert list(db.collection("assets").stream()) == []


def test_max_rows_with_assignments_in_one_transaction(client, admin_headers, db):
    """P10 상한(500행) 전부 배정 포함 — 자산 500 + 이력 500 + 카운터 쓰기가 한 트랜잭션에 들어가는지."""
    _member(client, admin_headers, "Z001", "김철수")
    rows = [{"품명": f"n{i}", "그룹": "A", "구매일": "2025-01-01", "사용자 사번": "Z001"} for i in range(ASSET_UPLOAD_MAX_ROWS)]
    r = _upload(client, admin_headers, xlsx({"SW": rows}))
    assert r.status_code == 201, r.text
    assert len(r.json()["asset_nos"]) == ASSET_UPLOAD_MAX_ROWS
    assert db.collection("asset_seq").document("SW-25").get().get("last_seq") == ASSET_UPLOAD_MAX_ROWS
